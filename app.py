import io
import os
import base64
import json
from datetime import date, datetime, timedelta

import gspread
import pandas as pd
import requests
import streamlit as st
from google.oauth2.service_account import Credentials
from PIL import Image, ImageDraw, ImageFont

def get_page_icon():
    candidates = [
        os.path.join(os.path.dirname(__file__), "logo.png"),
        os.path.join(os.path.dirname(__file__), "2.png"),
        "logo.png",
        "2.png",
    ]
    for p in candidates:
        if os.path.exists(p):
            try:
                return Image.open(p)
            except Exception:
                return p
    return "📦"

try:
    fav_icon = get_page_icon()
except Exception:
    fav_icon = "📦"

st.set_page_config(page_title="VIFEX - Quản lý đơn hàng", page_icon=fav_icon, layout="centered")

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive"
]

GREEN = "#15503F"
GREEN_BG = "#E2EDE8"
RED = "#D92B2B"
RED_BG = "#FDE8E8"
AMBER = "#D97706"
AMBER_BG = "#FEF3C7"
BLUE = "#2563EB"
BLUE_BG = "#EFF6FF"
PURPLE = "#7C3AED"
PURPLE_BG = "#F5F3FF"
TEAL = "#0D9488"
TEAL_BG = "#CCFBF1"
ORANGE = "#EA580C"
ORANGE_BG = "#FFEDD5"
GRAY_BG = "#F3F4F6"
GRAY_TEXT = "#4B5563"

ORDER_STATUSES = ["Lên đơn", "Gửi kho", "Đang giao", "Đã giao"]
VALID_ORDER_STATUSES = ["Gửi kho", "Đang giao", "Đang giao hàng", "Đã giao", "Đã giao hàng", "Đã nhận hàng"]

PAYMENT_STATUSES = [
    "Chưa thanh toán",
    "Chưa TT - Nháp",
    "Chưa TT - Đã VAT",
    "TT - Chưa VAT",
    "TT - Nháp VAT",
    "TT - không VAT",
    "TT - Đã VAT"
]

VAT_FOLDER_ID = "1HZiL99pNqV31u6Z8q5EqeoyjFJkPJFji"
DEFAULT_SCRIPT_URL = "https://script.google.com/macros/s/AKfycbyLfDcPZIs2QEshLL_uYSQ6-N4CnSbJhmorx3EI_28QZGd1EnNUEF9yrzh3Zx8M3bgqNw/exec"

def safe_str(v):
    if v is None or pd.isna(v):
        return ""
    s = str(v).strip()
    return "" if s.lower() == "nan" else s

def money(v):
    try:
        return f"{float(v):,.0f}đ".replace(",", ".")
    except (TypeError, ValueError):
        return "0đ"

def export_df_to_excel(df_dict):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            for sheet_name, df in df_dict.items():
                clean_sheet_name = str(sheet_name)[:31]
                df.to_excel(writer, sheet_name=clean_sheet_name, index=False)
                ws = writer.sheets[clean_sheet_name]

                # Freeze dòng tiêu đề trên cùng
                ws.freeze_panes = "A2"

                # Màu sắc nhận diện thương hiệu VIFEX
                header_fill = PatternFill(start_color="15503F", end_color="15503F", fill_type="solid")
                header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

                row_even_fill = PatternFill(start_color="F8FAF9", end_color="F8FAF9", fill_type="solid")
                row_odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

                border_thin = Side(border_style="thin", color="D1D5DB")
                cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

                font_regular = Font(name="Segoe UI", size=10)
                font_bold = Font(name="Segoe UI", size=10, bold=True)

                # 1. Định dạng hàng tiêu đề
                ws.row_dimensions[1].height = 28
                for col_idx in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_align
                    cell.border = cell_border

                # 2. Định dạng dữ liệu từng dòng
                num_rows = len(df)
                for r_idx in range(2, num_rows + 2):
                    ws.row_dimensions[r_idx].height = 21
                    fill_current = row_even_fill if r_idx % 2 == 0 else row_odd_fill
                    for c_idx in range(1, len(df.columns) + 1):
                        cell = ws.cell(row=r_idx, column=c_idx)
                        cell.border = cell_border
                        cell.fill = fill_current
                        cell.font = font_regular
                        col_name = str(df.columns[c_idx - 1])

                        if col_name in ["Mã đơn", "Ngày lên đơn", "Trạng thái thanh toán"]:
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        elif "Số lượng" in col_name:
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            try:
                                cell.value = float(cell.value)
                            except (ValueError, TypeError):
                                pass
                            cell.number_format = '#,##0'
                        elif "Tổng tiền" in col_name or "Doanh thu" in col_name:
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            try:
                                cell.value = float(cell.value)
                            except (ValueError, TypeError):
                                pass
                            cell.number_format = '#,##0 "₫"'
                        else:
                            cell.alignment = Alignment(horizontal="left", vertical="center")

                # 3. Dòng Tổng cộng nổi bật ở cuối bảng
                if num_rows > 0:
                    tot_row = num_rows + 2
                    ws.row_dimensions[tot_row].height = 25
                    tot_fill = PatternFill(start_color="E2EDE8", end_color="E2EDE8", fill_type="solid")
                    double_bottom = Side(border_style="double", color="15503F")
                    top_thin = Side(border_style="thin", color="15503F")
                    tot_border = Border(left=border_thin, right=border_thin, top=top_thin, bottom=double_bottom)

                    for c_idx in range(1, len(df.columns) + 1):
                        c_cell = ws.cell(row=tot_row, column=c_idx)
                        c_cell.fill = tot_fill
                        c_cell.font = font_bold
                        c_cell.border = tot_border
                        col_name = str(df.columns[c_idx - 1])
                        col_letter = get_column_letter(c_idx)

                        if c_idx == 1:
                            c_cell.value = "TỔNG CỘNG"
                            c_cell.alignment = Alignment(horizontal="center", vertical="center")
                        elif "Số lượng" in col_name:
                            c_cell.value = f"=SUM({col_letter}2:{col_letter}{tot_row-1})"
                            c_cell.alignment = Alignment(horizontal="right", vertical="center")
                            c_cell.number_format = '#,##0'
                        elif "Tổng tiền" in col_name or "Doanh thu" in col_name:
                            c_cell.value = f"=SUM({col_letter}2:{col_letter}{tot_row-1})"
                            c_cell.alignment = Alignment(horizontal="right", vertical="center")
                            c_cell.number_format = '#,##0 "₫"'

                # 4. Tự động co giãn vừa vặn độ rộng cột (Auto-fit width)
                for c_idx in range(1, len(df.columns) + 1):
                    col_letter = get_column_letter(c_idx)
                    max_len = 0
                    for cell in ws[col_letter]:
                        val_str = str(cell.value or "")
                        if str(val_str).startswith("="):
                            val_str = "123,456,789 ₫"
                        max_len = max(max_len, len(val_str))
                    ws.column_dimensions[col_letter].width = min(max(max_len + 5, 14), 55)

        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    except Exception as e:
        first_df = list(df_dict.values())[0] if df_dict else pd.DataFrame()
        csv_bytes = first_df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")
        return csv_bytes, "text/csv", "csv"

def get_logo_base64():
    candidates = ["logo.png", "2.png", os.path.join(os.path.dirname(__file__), "logo.png")]
    for p in candidates:
        if os.path.exists(p):
            with open(p, "rb") as img_f:
                return base64.b64encode(img_f.read()).decode()
    return None

def get_qr_bank_image():
    candidates = ["qr_bank.png", "qr.png", os.path.join(os.path.dirname(__file__), "qr_bank.png")]
    for p in candidates:
        if os.path.exists(p):
            try:
                return Image.open(p).convert("RGBA")
            except Exception:
                pass
    return None

st.markdown("""
<style>
header[data-testid="stHeader"] { background-color: transparent !important; z-index: 1 !important; }
.block-container { max-width: 960px !important; padding-top: 3.5rem !important; padding-bottom: 3rem !important; margin: 0 auto !important; }
.vifex-banner { background: #15503F; color: #ffffff; padding: 16px 22px; border-radius: 14px; margin-bottom: 14px; display: flex; justify-content: space-between; align-items: center; }
.vifex-banner .brand-tag { display: flex; align-items: center; gap: 6px; font-size: 13.5px; font-weight: 700; letter-spacing: 1px; }
.vifex-banner .brand-tag::before { content: ""; display: inline-block; width: 6px; height: 14px; background: #E53E3E; border-radius: 3px; }
.vifex-banner .main-title { font-size: 23px; font-weight: 700; margin-top: 2px; }
.vifex-banner-logo { background: #ffffff; width: 64px; height: 64px; border-radius: 12px; padding: 5px; display: flex; align-items: center; justify-content: center; }
.vifex-banner-logo img { max-width: 100%; max-height: 100%; object-fit: contain; }
div[data-testid="stVerticalBlockBorderWrapper"] { margin-bottom: 8px !important; border-radius: 12px !important; background: #ffffff !important; }
.order-code-compact { font-weight: 700; font-size: 15.5px; color: #111827; }
.order-cust-compact { font-size: 13px; color: #4b5563; margin-bottom: 3px; }
.order-value-compact { font-weight: 700; color: #15503F; font-size: 15px; }
.badge { display: inline-block; padding: 3px 8px; border-radius: 6px; font-size: 11.5px; font-weight: 600; text-align: center; }
.metric-box { background: #ffffff; border: 1px solid #edf0ed; border-radius: 12px; padding: 12px 14px; }
.metric-label { font-size: 12px; color: #6b7280; font-weight: 500; }
.metric-value { font-size: 19px; font-weight: 700; margin-top: 3px; }
.quick-order-card { background-color: #F8FAF9; border-left: 4px solid #15503F; border-radius: 10px; padding: 14px 18px; font-size: 13.8px; line-height: 1.7; margin-bottom: 10px; }
.quick-order-card .label { font-weight: 700; min-width: 105px; display: inline-block; }
.order-item-box { background-color: #ffffff; border: 1px solid #E5E7EB; border-radius: 10px; padding: 10px 14px; margin-bottom: 8px; }
.order-item-title { font-size: 13px; font-weight: 700; color: #15503F; }
.order-live-total-card { background-color: #E2EDE8; border: 1.5px solid #15503F; border-radius: 12px; padding: 12px 16px; display: flex; justify-content: space-between; align-items: center; margin: 12px 0; }
.order-live-total-title { font-size: 13.5px; font-weight: 700; color: #15503F; }
.order-live-total-value { font-size: 21px; font-weight: 800; color: #D92B2B; }
.card-wrapper { border-radius: 14px; padding: 8px 10px; min-height: 66px; display: flex; flex-direction: column; justify-content: center; }
.card-num { font-size: 22px; font-weight: 800; line-height: 1.1; }
.card-title { font-size: 12px; font-weight: 600; margin-top: 3px; }
div[class*="st-key-btn_filter_"] { margin-top: -65px; opacity: 0; }
div[class*="st-key-btn_filter_"] button { height: 65px !important; width: 100% !important; }
</style>
""", unsafe_allow_html=True)

def banner(title, subtitle=None, highlight_text=None):
    sub_html = f'<div class="sub-title">{subtitle}</div>' if subtitle else ""
    main_text = highlight_text if highlight_text else title
    logo_b64 = get_logo_base64()
    logo_html = f'<div class="vifex-banner-logo"><img src="data:image/png;base64,{logo_b64}" /></div>' if logo_b64 else '<div class="vifex-banner-logo" style="font-weight:800;color:#15503F;">VIFEX</div>'
    st.markdown(f'<div class="vifex-banner"><div><div class="brand-tag">VIFEX</div>{sub_html}<div class="main-title">{main_text}</div></div>{logo_html}</div>', unsafe_allow_html=True)

def order_status_badge_html(status):
    s = safe_str(status).lower()
    if "lên đơn" in s: return '<span class="badge" style="background:#EEF6FF;color:#1D4ED8;">Lên đơn</span>'
    if "gửi kho" in s: return '<span class="badge" style="background:#F5EEFF;color:#7C3AED;">Gửi kho</span>'
    if any(k in s for k in ["đang giao", "giao hàng"]): return '<span class="badge" style="background:#FFF6DA;color:#D97706;">Đang giao</span>'
    if any(k in s for k in ["đã giao", "đã nhận"]): return '<span class="badge" style="background:#E6F4EA;color:#065F46;">Đã giao</span>'
    return f'<span class="badge" style="background:{GRAY_BG};color:{GRAY_TEXT}">{status}</span>'

def payment_status_badge_html(status):
    s = safe_str(status).lower()
    if "chưa tt - nháp" in s: return '<span class="badge" style="background:#FFF1F2;color:#BE123C;">Chưa TT - Nháp</span>'
    if "đã vat" in s and "chưa" in s: return '<span class="badge" style="background:#FDF2F8;color:#DB2777;">Chưa TT - Đã VAT</span>'
    if "tt - chưa vat" in s: return '<span class="badge" style="background:#FEF3C7;color:#D97706;">TT - Chưa VAT</span>'
    if "tt - nháp vat" in s: return '<span class="badge" style="background:#FFF0E5;color:#EA580C;">TT - Nháp VAT</span>'
    if "không vat" in s: return '<span class="badge" style="background:#DCFAF4;color:#0D9488;">TT - không VAT</span>'
    if "tt - đã vat" in s: return '<span class="badge" style="background:#E5F0EC;color:#2D6A4F;">TT - Đã VAT</span>'
    if any(k in s for k in ["chưa thanh toán", "chưa tt"]): return '<span class="badge" style="background:#FFEBEF;color:#E11D48;">Chưa TT</span>'
    return f'<span class="badge" style="background:{GRAY_BG};color:{GRAY_TEXT}">{status}</span>'

@st.cache_resource
def get_credentials():
    return Credentials.from_service_account_info(st.secrets["gcp_service_account"], scopes=SCOPES)

@st.cache_resource
def get_client():
    return gspread.authorize(get_credentials())

@st.cache_resource
def get_spreadsheet():
    return get_client().open_by_key(st.secrets["SPREADSHEET_ID"])

def serial_to_date(v):
    if v is None or v == "": return None
    try:
        if isinstance(v, str): return pd.to_datetime(v, errors="coerce").date()
        return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))).date()
    except Exception: return None

SHEET_NAMES = ["San_pham", "Khach_hang", "Nhan_vien", "Lich_su_gia", "Don_hang", "Chi_tiet_don_hang"]

def read_all_sheets_batch(names):
    sh = get_spreadsheet()
    ranges = [f"{n}!A:Z" for n in names]
    resp = sh.values_batch_get(ranges, params={"valueRenderOption": "UNFORMATTED_VALUE"})
    result = {}
    for n, vr in zip(names, resp.get("valueRanges", [])):
        values = vr.get("values", [])
        if not values:
            result[n] = pd.DataFrame()
            continue
        header = [str(c).strip() if c is not None else "" for c in values[0]]
        if n == "Khach_hang" and (not header or header[0] in ["0", "", "none"]):
            header[0] = "Ma_KH"
        rows = values[1:]
        rows = [r + [None] * (len(header) - len(r)) for r in rows]
        df = pd.DataFrame(rows, columns=header)
        df = df.replace("", None)
        result[n] = df
    return result

def to_date_col(df, col):
    if col in df.columns: df[col] = df[col].apply(serial_to_date)
    return df

@st.cache_data(ttl=60)
def load_data():
    raw = read_all_sheets_batch(SHEET_NAMES)
    san_pham_df = raw["San_pham"]
    khach_hang_df = raw["Khach_hang"]
    nhan_vien_df = raw["Nhan_vien"]
    lich_su_gia_df = to_date_col(to_date_col(raw["Lich_su_gia"], "Ngay_bat_dau"), "Ngay_ket_thuc")
    don_hang_df = to_date_col(raw["Don_hang"], "Ngay_len_don")
    ctdh_df = raw["Chi_tiet_don_hang"]

    if not khach_hang_df.empty:
        if "Ma_KH" not in khach_hang_df.columns:
            khach_hang_df.rename(columns={khach_hang_df.columns[0]: "Ma_KH"}, inplace=True)
        if "Ten_NPP" not in khach_hang_df.columns and len(khach_hang_df.columns) > 1:
            khach_hang_df.rename(columns={khach_hang_df.columns[1]: "Ten_NPP"}, inplace=True)

    if "Trang_thai_TT" not in don_hang_df.columns:
        def map_pay(r):
            st_val = safe_str(r.get("Trang_thai"))
            for p in PAYMENT_STATUSES:
                if p.lower() in st_val.lower(): return p
            return "Chưa thanh toán"
        don_hang_df["Trang_thai_TT"] = don_hang_df.apply(map_pay, axis=1)
    else:
        don_hang_df["Trang_thai_TT"] = don_hang_df["Trang_thai_TT"].fillna("Chưa thanh toán")

    def norm_order(s):
        s = safe_str(s).lower()
        if "lên đơn" in s: return "Lên đơn"
        if "gửi kho" in s: return "Gửi kho"
        if any(k in s for k in ["đang giao", "giao hàng"]): return "Đang giao"
        if any(k in s for k in ["đã giao", "đã nhận"]): return "Đã giao"
        return "Lên đơn"
    
    don_hang_df["Trang_thai_Don"] = don_hang_df["Trang_thai"].apply(norm_order)

    for col in ["SL_dat", "Tang", "Don_gia_ap_dung", "Chiet_khau", "Thanh_tien", "San_luong_xuat_kho"]:
        if col in ctdh_df.columns:
            ctdh_df[col] = pd.to_numeric(ctdh_df[col], errors="coerce").fillna(0)

    return {
        "san_pham": san_pham_df, "khach_hang": khach_hang_df, "nhan_vien": nhan_vien_df,
        "lich_su_gia": lich_su_gia_df, "don_hang": don_hang_df, "ctdh": ctdh_df
    }

def get_ws(name): return get_spreadsheet().worksheet(name)
def refresh(): load_data.clear()

def next_code(ws, col_index, prefix, width):
    col = ws.col_values(col_index)[1:]
    nums = [int(str(v).replace(prefix, "").strip()) for v in col if str(v).replace(prefix, "").strip().isdigit()]
    return f"{prefix}{str((max(nums) + 1) if nums else 1).zfill(width)}"

def lookup_gia(ma_sp, ngay, lich_su_gia_df):
    df = lich_su_gia_df[lich_su_gia_df["Ma_SP"] == ma_sp]
    if df.empty: return 0
    valid = df[(df["Ngay_bat_dau"].notna()) & (df["Ngay_bat_dau"] <= ngay) & (df["Ngay_ket_thuc"].isna() | (df["Ngay_ket_thuc"] >= ngay))].sort_values("Ngay_bat_dau")
    return float(valid.iloc[-1]["Gia_ap_dung"]) if not valid.empty else 0

def find_row_by_code(ws, code, col_index=1):
    for i, v in enumerate(ws.col_values(col_index), start=1):
        if str(v).strip() == str(code).strip(): return i
    return None

def get_kh_dict(ma_kh_val, kh_df):
    if not ma_kh_val or kh_df.empty: return {}
    col_ma = "Ma_KH" if "Ma_KH" in kh_df.columns else kh_df.columns[0]
    matched = kh_df[kh_df[col_ma].astype(str).str.strip().str.lower() == str(ma_kh_val).strip().lower()]
    return matched.iloc[0].to_dict() if not matched.empty else {}

def update_order_both_statuses(ma_don, new_order_status, new_payment_status):
    today = (datetime.utcnow() + timedelta(hours=7)).strftime("%Y-%m-%d")
    don_hang_ws = get_ws("Don_hang")
    row = find_row_by_code(don_hang_ws, ma_don, col_index=1)
    if row:
        headers = don_hang_ws.row_values(1)
        col_pay = (headers.index("Trang_thai_TT") + 1) if "Trang_thai_TT" in headers else 11
        don_hang_ws.batch_update([
            {"range": gspread.utils.rowcol_to_a1(row, 5), "values": [[new_order_status]]},
            {"range": gspread.utils.rowcol_to_a1(row, 8), "values": [[today]]},
            {"range": gspread.utils.rowcol_to_a1(row, col_pay), "values": [[new_payment_status]]}
        ], value_input_option="USER_ENTERED")

    ctdh_ws = get_ws("Chi_tiet_don_hang")
    ma_don_col = ctdh_ws.col_values(2)
    updates = [{"range": gspread.utils.rowcol_to_a1(i, 10), "values": [[new_order_status]]} for i, v in enumerate(ma_don_col, start=1) if str(v).strip() == str(ma_don).strip()]
    if updates: ctdh_ws.batch_update(updates, value_input_option="USER_ENTERED")
    refresh()

def delete_order_completely(ma_don):
    try:
        ws = get_ws("Chi_tiet_don_hang")
        vals = ws.col_values(2)
        for i in range(len(vals), 0, -1):
            if str(vals[i-1]).strip() == str(ma_don).strip(): ws.delete_rows(i)
    except Exception: pass
    try:
        ws = get_ws("Don_hang")
        row = find_row_by_code(ws, ma_don, 1)
        if row: ws.delete_rows(row)
    except Exception: pass
    delete_vat_record(ma_don)
    refresh()

def get_upload_endpoint():
    try:
        u = st.secrets.get("DRIVE_UPLOAD_URL", "").strip()
        if u.startswith("http"): return u
    except Exception: pass
    return DEFAULT_SCRIPT_URL

def get_vat_sheet():
    sh = get_spreadsheet()
    try: return sh.worksheet("Hoa_don_VAT")
    except Exception:
        ws = sh.add_worksheet(title="Hoa_don_VAT", rows=100, cols=6)
        ws.append_row(["Ma_HD_VAT", "Ma_don", "Ma_KH", "Ten_file", "Ngay_tai_len", "Link_Drive"])
        return ws

def get_vat_link_from_sheet(ma_don):
    try:
        for r in get_vat_sheet().get_all_records():
            if str(r.get("Ma_don")).strip() == str(ma_don).strip():
                l = r.get("Link_Drive") or r.get("Link_file_PDF") or ""
                if str(l).startswith("http"): return l
    except Exception: pass
    return None

def upload_vat_directly_to_drive(ma_don, ma_kh, uploaded_file):
    url = get_upload_endpoint()
    _, ext = os.path.splitext(uploaded_file.name)
    payload = {
        "folderId": VAT_FOLDER_ID,
        "fileName": f"{ma_don}_VAT{ext.lower()}",
        "mimeType": uploaded_file.type or "application/pdf",
        "base64Data": base64.b64encode(uploaded_file.getvalue()).decode()
    }
    resp = requests.post(url, data=json.dumps(payload), headers={"Content-Type": "application/json"}, timeout=60).json()
    if resp.get("status") == "success":
        file_url = resp.get("fileUrl")
        ws = get_vat_sheet()
        row = find_row_by_code(ws, ma_don, 2)
        today_str = (datetime.utcnow() + timedelta(hours=7)).strftime("%Y-%m-%d %H:%M")
        if row:
            ws.update_cell(row, 4, uploaded_file.name); ws.update_cell(row, 5, today_str); ws.update_cell(row, 6, file_url)
        else:
            ws.append_row([next_code(ws, 1, "HDVAT", 4), ma_don, ma_kh, uploaded_file.name, today_str, file_url])
        return file_url
    raise Exception(resp.get("message", "Lỗi tải file"))

def delete_vat_record(ma_don):
    try:
        ws = get_vat_sheet()
        row = find_row_by_code(ws, ma_don, 2)
        if row: ws.delete_rows(row)
    except Exception: pass

@st.cache_resource
def get_font(size):
    candidates = [os.path.join(os.path.dirname(__file__), "fonts", "NotoSans.ttf"), "NotoSans.ttf"]
    for p in candidates:
        if os.path.exists(p): return ImageFont.truetype(p, size)
    raise FileNotFoundError("NotoSans font missing.")

def draw_bold(draw, pos, text, font, fill):
    x, y = pos
    draw.text((x, y), text, font=font, fill=fill)
    draw.text((x + 0.6, y), text, font=font, fill=fill)

def draw_wrapped_text(d, pos, text, font, fill, max_width, line_spacing=6):
    x, y = pos
    lines, curr = [], ""
    for w in text.split(" "):
        test = (curr + " " + w).strip()
        if (d.textbbox((0, 0), test, font=font)[2] - d.textbbox((0, 0), test, font=font)[0]) <= max_width:
            curr = test
        else:
            if curr: lines.append(curr)
            curr = w
    if curr: lines.append(curr)
    for l in lines:
        d.text((x, y), l, font=font, fill=fill)
        y += (d.textbbox((0, 0), l, font=font)[3] - d.textbbox((0, 0), l, font=font)[1]) + line_spacing
    return y

def generate_order_slip(ma_don, order_row, items_df, kh_dict):
    W, row_h = 860, 34
    f_title, f_h, f_n = get_font(26), get_font(16), get_font(14)
    H = 540 + row_h * (len(items_df) + 2) + 140
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, 80], fill=GREEN)
    draw_bold(d, (24, 16), "VIFEX", f_title, "white")
    d.text((24, 50), "PHIẾU XUẤT ĐƠN HÀNG (không phải hóa đơn VAT)", font=get_font(14), fill="white")

    y = 100
    ten_cty = safe_str(kh_dict.get("Ten_cong_ty_GPKD")) or safe_str(kh_dict.get("Ten_NPP"))
    ten_npp = safe_str(kh_dict.get("Ten_NPP"))
    dia_chi = safe_str(kh_dict.get("Dia_chi_giao_phu")) or safe_str(kh_dict.get("Dia_chi_GPKD"))
    sdt = safe_str(kh_dict.get("SDT_phu"))
    mst = safe_str(kh_dict.get("MST"))

    draw_bold(d, (24, y), f"Mã đơn: {ma_don}", f_h, "black")
    if mst:
        tw = d.textbbox((0, 0), f"MST: {mst}", font=f_n)[2] - d.textbbox((0, 0), f"MST: {mst}", font=f_n)[0]
        draw_bold(d, (W - 24 - tw, y + 1), f"MST: {mst}", f_n, "black")
    y += 26
    d.text((24, y), f"Ngày lên đơn: {order_row.get('Ngay_len_don')}", font=f_n, fill="black"); y += 24
    if ten_cty: y = draw_wrapped_text(d, (24, y), f"Khách hàng: {ten_cty}", f_n, "black", 810)
    if ten_npp: y = draw_wrapped_text(d, (24, y), f"NHÀ PHÂN PHỐI : {ten_npp}", f_n, "black", 810)
    if dia_chi: y = draw_wrapped_text(d, (24, y), f"Địa chỉ giao: {dia_chi}", f_n, "black", 810)
    if sdt: d.text((24, y), f"SĐT người nhận: {sdt}", font=f_n, fill="black"); y += 24
    y += 8; d.line([24, y, W - 24, y], fill="#ddd", width=1); y += 10

    cols_x = [24, 270, 370, 440, 520, 640, 750]
    for x, h in zip(cols_x, ["Sản phẩm", "Đơn giá", "SL đặt", "Tặng", "Tổng tiền", "Chiết khấu", "Thành tiền"]):
        draw_bold(d, (x, y), h, f_n, GREEN)
    y += row_h; d.line([24, y - 6, W - 24, y - 6], fill="#ddd", width=1)

    total = 0
    for _, r in items_df.iterrows():
        thanh_tien = float(r["Thanh_tien"])
        total += thanh_tien
        d.text((cols_x[0], y), str(r["Ten_SP"])[:26], font=f_n, fill="black")
        d.text((cols_x[1], y), money(r["Don_gia_ap_dung"]), font=f_n, fill="black")
        d.text((cols_x[2], y), str(int(r["SL_dat"])), font=f_n, fill="black")
        d.text((cols_x[3], y), str(int(r["Tang"])), font=f_n, fill="black")
        d.text((cols_x[4], y), money(float(r["SL_dat"]) * float(r["Don_gia_ap_dung"])), font=f_n, fill="black")
        d.text((cols_x[5], y), money(r["Chiet_khau"]), font=f_n, fill="black")
        d.text((cols_x[6], y), money(thanh_tien), font=f_n, fill="black")
        y += row_h

    d.line([24, y, W - 24, y], fill="#ddd", width=1); y += 14
    draw_bold(d, (cols_x[5], y), "TỔNG CỘNG:", f_h, RED)
    draw_bold(d, (cols_x[6], y), money(total), f_h, RED)
    y += 32
    d.text((24, y), f"Hình thức thanh toán: {order_row.get('Hinh_thuc_thanh_toan', '')}", font=f_n, fill="black"); y += 22
    if safe_str(order_row.get("Ghi_chu_thanh_toan")):
        y = draw_wrapped_text(d, (24, y), f"Ghi chú: {order_row['Ghi_chu_thanh_toan']}", f_n, "black", 810)

    y += 10; d.line([24, y, W - 24, y], fill="#eee", width=1); y += 14
    qr = get_qr_bank_image()
    if qr:
        img.paste(qr.resize((120, 120)), (24, y))
        draw_bold(d, (162, y + 10), "THÔNG TIN THANH TOÁN CHUYỂN KHOẢN:", f_h, GREEN)
        draw_bold(d, (162, y + 36), "STK Vietinbank: 116003017106", f_n, "black")
        d.text((162, y + 58), "Chủ TK: Công ty TNHH VIFEX", font=f_n, fill="black")
        y += 135
    else:
        draw_bold(d, (24, y), "STK Vietinbank: 116003017106 - Công ty TNHH VIFEX", f_n, GREEN); y += 28

    buf = io.BytesIO()
    img.crop((0, 0, W, max(y + 20, 320))).save(buf, format="PNG")
    return buf.getvalue()

data = load_data()
san_pham_df = data["san_pham"]
khach_hang_df = data["khach_hang"]
nhan_vien_df = data["nhan_vien"]
lich_su_gia_df = data["lich_su_gia"]
don_hang_df = data["don_hang"]
ctdh_df = data["ctdh"]

if khach_hang_df.empty or san_pham_df.empty:
    st.warning("Chưa có dữ liệu Khách hàng hoặc Sản phẩm."); st.stop()

merged = pd.DataFrame()
if not ctdh_df.empty and not don_hang_df.empty:
    dh_cols = [c for c in ["Ma_don", "Ngay_len_don", "Ma_KH", "Sale_phu_trach", "Trang_thai", "Trang_thai_Don", "Trang_thai_TT"] if c in don_hang_df.columns]
    merged = ctdh_df.merge(don_hang_df[dh_cols], on="Ma_don", how="left")

def order_total_quantity(ma_don):
    if ctdh_df.empty: return 0
    items = ctdh_df[ctdh_df["Ma_don"] == ma_don]
    return int(items.get("SL_dat", 0).sum() + items.get("Tang", 0).sum())

def order_categories(ma_don):
    if ctdh_df.empty or san_pham_df.empty: return ""
    items = ctdh_df[ctdh_df["Ma_don"] == ma_don]
    if items.empty: return ""
    col_dm = next((c for c in san_pham_df.columns if "nhom" in str(c).lower() or "danh_muc" in str(c).lower()), "Nhom_danh_muc")
    if col_dm in san_pham_df.columns:
        sp_m = items.merge(san_pham_df[["Ma_SP", col_dm]], on="Ma_SP", how="left")
        cats = [str(c).strip() for c in sp_m[col_dm].dropna().unique() if str(c).strip()]
        return ", ".join(cats) if cats else ""
    return ""

def order_total(ma_don):
    return merged.loc[merged["Ma_don"] == ma_don, "Thanh_tien"].sum() if not merged.empty else 0

if "selected_order" not in st.session_state: st.session_state.selected_order = None
if "nav" not in st.session_state: st.session_state.nav = "🏠 Trang chủ"
if "order_items_count" not in st.session_state: st.session_state.order_items_count = 1
if "order_form_version" not in st.session_state: st.session_state.order_form_version = 0
if "home_filter_category" not in st.session_state: st.session_state.home_filter_category = None
if "home_filter_value" not in st.session_state: st.session_state.home_filter_value = None

NAV_OPTIONS = ["Trang chủ", "Đơn hàng", "Lên đơn", "Lương Sale", "Dashboard", "Khách hàng"]
NAV_ICONS = {"Trang chủ": "🏠", "Đơn hàng": "📦", "Lên đơn": "➕", "Lương Sale": "💰", "Dashboard": "📊", "Khách hàng": "👥"}

with st.container(key="vifex_nav"):
    r1, r2 = st.columns(3), st.columns(3)
    for idx, opt in enumerate(NAV_OPTIONS[:3]):
        if r1[idx].button(f"{NAV_ICONS[opt]} {opt}", key=f"nav_{opt}", type="primary" if st.session_state.nav == f"{NAV_ICONS[opt]} {opt}" else "secondary", use_container_width=True):
            st.session_state.nav = f"{NAV_ICONS[opt]} {opt}"; st.rerun()
    for idx, opt in enumerate(NAV_OPTIONS[3:]):
        if r2[idx].button(f"{NAV_ICONS[opt]} {opt}", key=f"nav_{opt}", type="primary" if st.session_state.nav == f"{NAV_ICONS[opt]} {opt}" else "secondary", use_container_width=True):
            st.session_state.nav = f"{NAV_ICONS[opt]} {opt}"; st.rerun()

nav = st.session_state.nav

def render_order_detail_inline(ma_don):
    order_rows = don_hang_df[don_hang_df["Ma_don"] == ma_don]
    if order_rows.empty: return
    order_row = order_rows.iloc[0]
    kh_dict = get_kh_dict(safe_str(order_row.get("Ma_KH")), khach_hang_df)
    items = ctdh_df[ctdh_df["Ma_don"] == ma_don].merge(san_pham_df[["Ma_SP", "Ten_SP", "Nhom_danh_muc"]], on="Ma_SP", how="left")

    ten_cty = safe_str(kh_dict.get("Ten_cong_ty_GPKD")) or safe_str(kh_dict.get("Ten_NPP"))
    ten_npp = safe_str(kh_dict.get("Ten_NPP")) or safe_str(order_row.get("Ma_KH"))
    dia_chi = safe_str(kh_dict.get("Dia_chi_giao_phu")) or safe_str(kh_dict.get("Dia_chi_GPKD"))
    sdt = safe_str(kh_dict.get("SDT_phu"))
    ten_nhan = safe_str(kh_dict.get("Ten_nguoi_nhan_phu"))

    with st.container(border=True):
        st.markdown(f"#### Chi tiết đơn: `{ma_don}`")
        cb1, cb2 = st.columns(2)
        cb1.markdown(f"**Giao hàng:** {order_status_badge_html(order_row.get('Trang_thai_Don'))}", unsafe_allow_html=True)
        cb2.markdown(f"**Thanh toán:** {payment_status_badge_html(order_row.get('Trang_thai_TT'))}", unsafe_allow_html=True)
        st.write(f"**Khách hàng:** {ten_cty}")
        st.write(f"**NHÀ PHÂN PHỐI :** {ten_npp}")
        st.write(f"**Ngày lên đơn:** {order_row.get('Ngay_len_don')}")
        st.write(f"**Hình thức thanh toán:** {order_row.get('Hinh_thuc_thanh_toan', '')}")
        if safe_str(order_row.get("Ghi_chu_thanh_toan")): st.write(f"**Ghi chú:** {order_row['Ghi_chu_thanh_toan']}")

        nhom_str = ", ".join(items["Nhom_danh_muc"].dropna().unique().tolist()) or "Hàng hóa"
        nguoi_nhan_str = " - ".join([p for p in [ten_nhan, sdt, dia_chi] if p]) or "Chưa có thông tin nhận"
        with st.expander("📋 **Xem nhanh thông tin gửi hàng (dạng chữ):**", expanded=True):
            st.markdown(f"""
            <div class="quick-order-card">
                <div><span class="label">Đặt hàng:</span> {nhom_str}</div>
                <div><span class="label">NPP:</span> {ten_npp}</div>
                <div><span class="label">Ngày lên đơn:</span> {order_row.get('Ngay_len_don')}</div>
                <div><span class="label">Người nhận:</span> {nguoi_nhan_str}</div>
                <div><span class="label">Số lượng:</span> {int(items.get('SL_dat', 0).sum() + items.get('Tang', 0).sum())}</div>
                <div><span class="label">Ghi chú:</span> {safe_str(order_row.get('Ghi_chu_thanh_toan')) or 'Không có'}</div>
            </div>""", unsafe_allow_html=True)

        show_df = pd.DataFrame({
            "Sản phẩm": items["Ten_SP"], "Đơn giá": items["Don_gia_ap_dung"].apply(money),
            "SL đặt": items["SL_dat"].astype(int), "Tặng": items["Tang"].astype(int),
            "Tổng tiền": (items["SL_dat"] * items["Don_gia_ap_dung"]).apply(money),
            "Chiết khấu": items["Chiet_khau"].apply(money), "Thành tiền": items["Thanh_tien"].apply(money)
        })
        st.dataframe(show_df, hide_index=True, use_container_width=True)
        st.markdown(f"<div style='font-size:16px;font-weight:700;color:{GREEN};text-align:right;'>Tổng cộng: {money(items['Thanh_tien'].sum())}</div>", unsafe_allow_html=True)

        st.markdown("<div style='font-size:13px;font-weight:700;color:#15503F;margin:8px 0 4px 0;'>CẬP NHẬT TIẾN ĐỘ ĐƠN HÀNG & THANH TOÁN</div>", unsafe_allow_html=True)
        cst1, cst2, cst3 = st.columns([1.5, 1.5, 1])
        with cst1: new_order_st = st.selectbox("📦 Tiến độ giao hàng", ORDER_STATUSES, index=ORDER_STATUSES.index(order_row.get("Trang_thai_Don", "Lên đơn")), key=f"sel_o_{ma_don}")
        with cst2: new_pay_st = st.selectbox("💳 Thanh toán & VAT", PAYMENT_STATUSES, index=PAYMENT_STATUSES.index(order_row.get("Trang_thai_TT", "Chưa thanh toán")), key=f"sel_p_{ma_don}")
        with cst3:
            st.write(""); st.write("")
            if st.button("💾 Lưu thay đổi", key=f"save_{ma_don}", type="primary", use_container_width=True):
                update_order_both_statuses(ma_don, new_order_st, new_pay_st); st.success("Đã cập nhật!"); st.rerun()

        ca1, ca2 = st.columns(2)
        with ca1:
            preview_key = f"show_preview_{ma_don}"
            if preview_key not in st.session_state: st.session_state[preview_key] = False
            if st.button("🙈 Đóng phiếu xuất" if st.session_state[preview_key] else "🖼️ Xem Phiếu xuất đơn", key=f"btn_p_{ma_don}", use_container_width=True):
                st.session_state[preview_key] = not st.session_state[preview_key]; st.rerun()
        with ca2:
            vat_url = get_vat_link_from_sheet(ma_don)
            if vat_url: st.link_button("📄 Mở File Hóa đơn VAT (Drive)", url=vat_url, use_container_width=True)
            else: st.button("☁️ Chưa có File VAT", disabled=True, use_container_width=True)

        if st.session_state.get(f"show_preview_{ma_don}", False):
            png = generate_order_slip(ma_don, order_row, items, kh_dict)
            st.image(png, use_container_width=True)
            st.download_button("📥 Tải ảnh phiếu xuất (PNG)", data=png, file_name=f"{ma_don}_phieu_xuat.png", mime="image/png", key=f"dl_{ma_don}", use_container_width=True)

        ce1, ce2 = st.columns(2)
        with ce1:
            with st.expander("☁️ **Quản lý Hóa đơn VAT**"):
                vat_url = get_vat_link_from_sheet(ma_don)
                if vat_url:
                    st.info("✅ Đã có file trên Google Drive."); st.markdown(f"👉 [Mở xem file]({vat_url})")
                    if st.button("🗑️ Xóa hóa đơn này", key=f"del_vat_{ma_don}"):
                        delete_vat_record(ma_don); st.success("Đã xóa file!"); st.rerun()
                up = st.file_uploader(f"Chọn file cho đơn {ma_don}", type=["pdf", "png", "jpg"], key=f"up_vat_{ma_don}")
                if up and st.button("⬆️ Tải lên Google Drive", key=f"save_vat_{ma_don}", type="primary"):
                    upload_vat_directly_to_drive(ma_don, safe_str(order_row.get("Ma_KH")), up); st.success("Đã tải lên!"); st.rerun()
        with ce2:
            with st.expander("⚠️ **Xóa đơn hàng này**"):
                if st.checkbox(f"Xác nhận xóa {ma_don}", key=f"chk_del_{ma_don}") and st.button("🗑 Xóa vĩnh viễn", key=f"cfm_del_{ma_don}", type="primary"):
                    delete_order_completely(ma_don); st.session_state.selected_order = None; st.success(f"Đã xóa đơn {ma_don}!"); st.rerun()

        if st.button("← Đóng xem chi tiết", key=f"close_{ma_don}", use_container_width=True):
            st.session_state.selected_order = None; st.rerun()

if nav == "🏠 Trang chủ":
    pending_delivery = don_hang_df[don_hang_df["Trang_thai_Don"].isin(["Lên đơn", "Gửi kho", "Đang giao"])] if not don_hang_df.empty else pd.DataFrame()
    banner("Trang chủ", subtitle="Xin chào, Coco", highlight_text=f"{len(pending_delivery)} đơn đang vận hành")

    if don_hang_df.empty:
        st.info("Chưa có đơn hàng nào.")
    else:
        counts_order = don_hang_df["Trang_thai_Don"].value_counts()
        counts_pay = don_hang_df["Trang_thai_TT"].value_counts()

        st.markdown("<div style='font-size:13px;font-weight:700;color:#15503F;margin-bottom:6px;'>📦 TIẾN ĐỘ VẬN HÀNH ĐƠN HÀNG (Nhấn ô để lọc danh sách)</div>", unsafe_allow_html=True)
        c1, c2, c3, c4 = st.columns(4)
        for name, cnt, bg, bdr, num, col_o in [
            ("Lên đơn", int(counts_order.get('Lên đơn', 0)), "#EEF6FF", "#B9D7FE", "#1D4ED8", c1),
            ("Gửi kho", int(counts_order.get('Gửi kho', 0)), "#F5EEFF", "#E0CFFE", "#7C3AED", c2),
            ("Đang giao", int(counts_order.get('Đang giao', 0)), "#FFF6DA", "#FDE08B", "#D97706", c3),
            ("Đã giao", int(counts_order.get('Đã giao', 0)), "#E6F4EA", "#BBE1C7", "#065F46", c4)
        ]:
            act = (st.session_state.home_filter_category == "order" and st.session_state.home_filter_value == name)
            with col_o:
                st.markdown(f'<div class="card-wrapper" style="background:{bg};border:{3 if act else 1.5}px solid {num if act else bdr};"><div class="card-num" style="color:{num};">{cnt}</div><div class="card-title">{name}{" ✔️" if act else ""}</div></div>', unsafe_allow_html=True)
                with st.container(key=f"btn_filter_order_{name}"):
                    if st.button(name, key=f"btn_act_o_{name}", use_container_width=True):
                        st.session_state.home_filter_category = None if act else "order"
                        st.session_state.home_filter_value = None if act else name
                        st.session_state.selected_order = None; st.rerun()

        st.write("")
        st.markdown("<div style='font-size:13px;font-weight:700;color:#15503F;margin-bottom:6px;'>💳 TIẾN ĐỘ THANH TOÁN & HÓA ĐƠN VAT (Nhấn ô để lọc danh sách)</div>", unsafe_allow_html=True)
        pr1 = st.columns(3)
        for full, short, cnt, bg, bdr, num, col_o in [
            ("Chưa thanh toán", "Chưa TT", int(counts_pay.get('Chưa thanh toán', 0)), "#FFEBEF", "#FECDD3", "#E11D48", pr1[0]),
            ("Chưa TT - Nháp", "Chưa TT - Nháp", int(counts_pay.get('Chưa TT - Nháp', 0)), "#FFF1F2", "#FECDD3", "#BE123C", pr1[1]),
            ("Chưa TT - Đã VAT", "Chưa TT - Đã VAT", int(counts_pay.get('Chưa TT - Đã VAT', 0)), "#FDF2F8", "#FBCFE8", "#DB2777", pr1[2])
        ]:
            act = (st.session_state.home_filter_category == "pay" and st.session_state.home_filter_value == full)
            with col_o:
                st.markdown(f'<div class="card-wrapper" style="background:{bg};border:{3 if act else 1.5}px solid {num if act else bdr};"><div class="card-num" style="color:{num};">{cnt}</div><div class="card-title">{short}{" ✔️" if act else ""}</div></div>', unsafe_allow_html=True)
                with st.container(key=f"btn_filter_pay_{short}"):
                    if st.button(short, key=f"btn_act_p_{short}", use_container_width=True):
                        st.session_state.home_filter_category = None if act else "pay"
                        st.session_state.home_filter_value = None if act else full
                        st.session_state.selected_order = None; st.rerun()

        st.write("")
        pr2 = st.columns(4)
        for full, short, cnt, bg, bdr, num, col_o in [
            ("TT - Chưa VAT", "TT - Chưa VAT", int(counts_pay.get('TT - Chưa VAT', 0)), "#FEF3C7", "#FDE68A", "#D97706", pr2[0]),
            ("TT - Nháp VAT", "TT - Nháp VAT", int(counts_pay.get('TT - Nháp VAT', 0)), "#FFF0E5", "#FED7AA", "#EA580C", pr2[1]),
            ("TT - không VAT", "TT - không VAT", int(counts_pay.get('TT - không VAT', 0)), "#DCFAF4", "#99F6E4", "#0D9488", pr2[2]),
            ("TT - Đã VAT", "TT - Đã VAT", int(counts_pay.get('TT - Đã VAT', 0)), "#E5F0EC", "#C4E3D7", "#2D6A4F", pr2[3])
        ]:
            act = (st.session_state.home_filter_category == "pay" and st.session_state.home_filter_value == full)
            with col_o:
                st.markdown(f'<div class="card-wrapper" style="background:{bg};border:{3 if act else 1.5}px solid {num if act else bdr};"><div class="card-num" style="color:{num};">{cnt}</div><div class="card-title">{short}{" ✔️" if act else ""}</div></div>', unsafe_allow_html=True)
                with st.container(key=f"btn_filter_pay_{short}"):
                    if st.button(short, key=f"btn_act_p_{short}", use_container_width=True):
                        st.session_state.home_filter_category = None if act else "pay"
                        st.session_state.home_filter_value = None if act else full
                        st.session_state.selected_order = None; st.rerun()

        st.write("")
        disp_df = don_hang_df.copy()
        list_title = "ĐƠN HÀNG CẦN XỬ LÝ GẦN ĐÂY"
        if st.session_state.home_filter_category == "order":
            disp_df = disp_df[disp_df["Trang_thai_Don"] == st.session_state.home_filter_value]
            list_title = f"DANH SÁCH ĐƠN: {st.session_state.home_filter_value.upper()} ({len(disp_df)} đơn)"
        elif st.session_state.home_filter_category == "pay":
            disp_df = disp_df[disp_df["Trang_thai_TT"] == st.session_state.home_filter_value]
            list_title = f"DANH SÁCH ĐƠN: {st.session_state.home_filter_value.upper()} ({len(disp_df)} đơn)"
        else:
            disp_df = disp_df.head(6)

        ch1, ch2 = st.columns([3, 1])
        ch1.markdown(f"<div style='font-size:14px;font-weight:700;color:#15503F;'>📋 {list_title}</div>", unsafe_allow_html=True)
        if st.session_state.home_filter_category and ch2.button("✖ Bỏ lọc", key="btn_clear_hf", use_container_width=True):
            st.session_state.home_filter_category = st.session_state.home_filter_value = st.session_state.selected_order = None; st.rerun()

        disp_df = disp_df.sort_values("Ma_don", ascending=False)
        for _, r in disp_df.iterrows():
            kh_info = get_kh_dict(safe_str(r.get("Ma_KH")), khach_hang_df)
            ten_npp_disp = safe_str(kh_info.get("Ten_NPP")) or safe_str(r.get("Ma_KH")) or "Chưa có tên NPP"
            is_exp = (st.session_state.selected_order == r["Ma_don"])
            with st.container(border=True):
                cl, cr = st.columns([2.6, 1.4])
                cl.markdown(f'<div class="order-code-compact">{r["Ma_don"]}</div><div class="order-cust-compact">NPP {ten_npp_disp}</div><div class="order-value-compact">{money(order_total(r["Ma_don"]))}</div>', unsafe_allow_html=True)
                cr.markdown(f'<div style="text-align:right;margin-bottom:4px;">{order_status_badge_html(r.get("Trang_thai_Don"))} {payment_status_badge_html(r.get("Trang_thai_TT"))}</div>', unsafe_allow_html=True)
                if cr.button("Đóng chi tiết" if is_exp else "Xem chi tiết", key=f"hview_{r['Ma_don']}", use_container_width=True):
                    st.session_state.selected_order = None if is_exp else r["Ma_don"]; st.rerun()
            if is_exp: render_order_detail_inline(r["Ma_don"])

    if st.button("+ Tạo đơn mới", key="home_new_o", type="primary", use_container_width=True):
        st.session_state.nav = "➕ Lên đơn"; st.rerun()

elif nav == "📦 Đơn hàng":
    banner("Danh sách đơn hàng")
    time_options = ["Tất cả"]
    if not don_hang_df.empty:
        valid_dates = pd.to_datetime(don_hang_df["Ngay_len_don"].dropna(), errors="coerce").dropna()
        for m in valid_dates.dt.to_period("M").drop_duplicates().sort_values(ascending=False): time_options.append(f"Tháng {m.month:02d}/{m.year}")
        for y in valid_dates.dt.year.drop_duplicates().sort_values(ascending=False):
            if f"Năm {y}" not in time_options: time_options.append(f"Năm {y}")

    list_npp = ["Tất cả"] + sorted([str(x).strip() for x in khach_hang_df["Ten_NPP"].dropna().unique() if str(x).strip()])
    f1, f2, f3 = st.columns(3)
    filter_order_st = f1.multiselect("Lọc tiến độ giao hàng", ORDER_STATUSES, default=[], placeholder="Tất cả trạng thái", key="f_ost")
    filter_time = f2.selectbox("Lọc theo thời gian", time_options, key="f_time")
    filter_npp = f3.selectbox("Lọc theo Nhà phân phối", list_npp, key="f_npp")

    view_df = don_hang_df.copy()
    if filter_order_st: view_df = view_df[view_df["Trang_thai_Don"].isin(filter_order_st)]
    if filter_time != "Tất cả" and not view_df.empty:
        v_dates = pd.to_datetime(view_df["Ngay_len_don"], errors="coerce")
        if filter_time.startswith("Tháng "):
            m_v, y_v = map(int, filter_time.replace("Tháng ", "").strip().split("/"))
            view_df = view_df[(v_dates.dt.month == m_v) & (v_dates.dt.year == y_v)]
        elif filter_time.startswith("Năm "):
            view_df = view_df[v_dates.dt.year == int(filter_time.replace("Năm ", "").strip())]
    if filter_npp != "Tất cả" and not view_df.empty:
        matching_kh_ids = khach_hang_df[khach_hang_df["Ten_NPP"] == filter_npp]["Ma_KH"].dropna().tolist()
        view_df = view_df[view_df["Ma_KH"].isin(matching_kh_ids)]

    view_df = view_df.sort_values("Ma_don", ascending=False)
    if not view_df.empty:
        exp_orders = view_df.copy().merge(khach_hang_df[["Ma_KH", "Ten_NPP"]], on="Ma_KH", how="left")
        if "Ma_NV" in nhan_vien_df.columns and "Ten_NV" in nhan_vien_df.columns:
            exp_orders = exp_orders.merge(nhan_vien_df[["Ma_NV", "Ten_NV"]], left_on="Sale_phu_trach", right_on="Ma_NV", how="left")
            sale_vals = exp_orders["Ten_NV"].fillna(exp_orders["Sale_phu_trach"]).fillna("")
        else: sale_vals = exp_orders["Sale_phu_trach"].fillna("")

        out_df = pd.DataFrame({
            "Mã đơn": exp_orders["Ma_don"],
            "Ngày lên đơn": exp_orders["Ngay_len_don"].astype(str),
            "Nhà phân phối": exp_orders["Ten_NPP"].fillna("Chưa có NPP"),
            "Nhân viên Sale": sale_vals,
            "Nhóm danh mục": exp_orders["Ma_don"].apply(order_categories),
            "Số lượng": exp_orders["Ma_don"].apply(order_total_quantity),
            "Tổng tiền sau CK (VNĐ)": exp_orders["Ma_don"].apply(order_total),
            "Trạng thái thanh toán": exp_orders["Trang_thai_TT"].fillna("Chưa thanh toán"),
            "Ghi chú": exp_orders["Ghi_chu_thanh_toan"].fillna("")
        })
        b_data, m_type, ext = export_df_to_excel({"Danh sách đơn hàng": out_df})
        cd1, cd2 = st.columns([2.5, 1.5])
        cd1.caption(f"Đang hiển thị **{len(view_df)}** đơn hàng.")
        cd2.download_button("📥 Tải dữ liệu (Excel)", data=b_data, file_name=f"VIFEX_DonHang_{date.today().strftime('%Y%m%d')}.{ext}", mime=m_type, key="btn_dl_excel", use_container_width=True)

    for _, r in view_df.iterrows():
        kh_info = get_kh_dict(safe_str(r.get("Ma_KH")), khach_hang_df)
        ten_npp_disp = safe_str(kh_info.get("Ten_NPP")) or safe_str(r.get("Ma_KH")) or "Chưa có tên NPP"
        is_exp = (st.session_state.selected_order == r["Ma_don"])
        with st.container(border=True):
            cl, cr = st.columns([2.6, 1.4])
            cl.markdown(f'<div class="order-code-compact">{r["Ma_don"]}</div><div class="order-cust-compact">NPP {ten_npp_disp}</div><div class="order-value-compact">{money(order_total(r["Ma_don"]))}</div>', unsafe_allow_html=True)
            cr.markdown(f'<div style="text-align:right;margin-bottom:4px;">{order_status_badge_html(r.get("Trang_thai_Don"))} {payment_status_badge_html(r.get("Trang_thai_TT"))}</div>', unsafe_allow_html=True)
            if cr.button("Đóng chi tiết" if is_exp else "Xem chi tiết", key=f"lview_{r['Ma_don']}", use_container_width=True):
                st.session_state.selected_order = None if is_exp else r["Ma_don"]; st.rerun()
        if is_exp: render_order_detail_inline(r["Ma_don"])

elif nav == "➕ Lên đơn":
    banner("Lên đơn hàng")
    v = st.session_state.order_form_version
    ten_npp = st.selectbox("Khách hàng (NPP)", khach_hang_df["Ten_NPP"].dropna().tolist(), key=f"fnpp_{v}")
    kh_matches = khach_hang_df[khach_hang_df["Ten_NPP"] == ten_npp]
    kh_row = kh_matches.iloc[0] if not kh_matches.empty else {}
    ma_kh = safe_str(kh_row.get("Ma_KH"))
    sale_pt = safe_str(kh_row.get("Sale_phu_trach"))
    st.caption(f"Sale phụ trách: **{sale_pt}**")

    if "just_created_order" in st.session_state:
        msg, money_val = st.session_state.just_created_order
        st.success(f"Đã tạo đơn **{msg}** — tổng giá trị: **{money(money_val)}**")
        del st.session_state.just_created_order

    with st.container(border=True):
        cd1, cd2 = st.columns(2)
        ngay_len_don = cd1.date_input("Ngày lên đơn", value=date.today(), key=f"fdate_{v}")
        hinh_thuc_tt = cd2.selectbox("Hình thức TT", ["Tiền mặt", "Chuyển khoản", "Công nợ", "Khác"], key=f"fhttt_{v}")
        ghi_chu_tt = st.text_input("Ghi chú thanh toán", "", key=f"fnote_{v}")

        st.markdown("<div style='font-size:14px;font-weight:700;color:#15503F;margin:10px 0 6px 0;'>DANH SÁCH SẢN PHẨM</div>", unsafe_allow_html=True)
        col_tt_sp = next((c for c in san_pham_df.columns if "trang_thai" in str(c).lower() or "trạng thái" in str(c).lower()), None)
        if col_tt_sp:
            active_sp = san_pham_df[~san_pham_df[col_tt_sp].astype(str).str.strip().str.lower().str.contains("ngừng|ngung|dừng|dung|khoá|khóa", na=False)]
            ten_sp_list = active_sp["Ten_SP"].dropna().tolist() or san_pham_df["Ten_SP"].dropna().tolist()
        else:
            ten_sp_list = san_pham_df["Ten_SP"].dropna().tolist()

        line_items, est_total = [], 0.0
        for i in range(st.session_state.order_items_count):
            with st.container():
                st.markdown(f'<div class="order-item-box"><div class="order-item-title">Sản phẩm #{i+1}</div></div>', unsafe_allow_html=True)
                ten_sp = st.selectbox(f"Chọn SP #{i+1}", ten_sp_list, key=f"sp_{v}_{i}", label_visibility="collapsed")
                c1, c2, c3 = st.columns(3)
                sl_dat = c1.number_input("SL đặt", min_value=0, value=0, key=f"sl_{v}_{i}")
                tang = c2.number_input("Tặng", min_value=0.0, value=0.0, step=0.1, key=f"t_{v}_{i}")
                chiet_khau = c3.number_input("CK (đ)", min_value=0, value=0, step=10000, key=f"ck_{v}_{i}")
                if chiet_khau > 0: c3.caption(f"↳ {money(chiet_khau)}")
                sp_row = san_pham_df[san_pham_df["Ten_SP"] == ten_sp]
                cur_price = lookup_gia(safe_str(sp_row.iloc[0]["Ma_SP"]), ngay_len_don, lich_su_gia_df) if not sp_row.empty else 0.0
                est_total += max(0.0, (float(sl_dat) * cur_price) - float(chiet_khau))
                line_items.append((ten_sp, sl_dat, tang, chiet_khau))

        cadd, crem = st.columns(2)
        if cadd.button("➕ Thêm sản phẩm", key="badd", use_container_width=True):
            st.session_state.order_items_count += 1; st.rerun()
        if crem.button("➖ Bớt sản phẩm", key="brem", use_container_width=True) and st.session_state.order_items_count > 1:
            st.session_state.order_items_count -= 1; st.rerun()

        st.markdown(f'<div class="order-live-total-card"><span class="order-live-total-title">TỔNG GIÁ TRỊ ĐƠN HÀNG (DỰ KIẾN):</span><span class="order-live-total-value">{money(est_total)}</span></div>', unsafe_allow_html=True)
        if st.button("✅ Tạo đơn hàng", use_container_width=True, type="primary"):
            valid_items = [li for li in line_items if li[1] > 0]
            if not valid_items: st.error("Cần ít nhất 1 dòng có SL đặt > 0."); st.stop()
            dh_ws, ctdh_ws = get_ws("Don_hang"), get_ws("Chi_tiet_don_hang")
            ma_don = next_code(dh_ws, 1, "DH", 5)
            dh_ws.append_row([str(ma_don), str(ngay_len_don.strftime("%Y-%m-%d")), str(ma_kh), str(sale_pt), "Lên đơn", str(hinh_thuc_tt), str(ghi_chu_tt), str(ngay_len_don.strftime("%Y-%m-%d")), int(ngay_len_don.month), int(ngay_len_don.year), "Chưa thanh toán"], value_input_option="USER_ENTERED")
            tot = 0
            for sp_name, sl, tg, ck in valid_items:
                sp_m = safe_str(san_pham_df[san_pham_df["Ten_SP"] == sp_name].iloc[0]["Ma_SP"])
                pr = lookup_gia(sp_m, ngay_len_don, lich_su_gia_df)
                tt = float(sl * pr - ck)
                tot += tt
                ctdh_ws.append_row([next_code(ctdh_ws, 1, "CT", 5), str(ma_don), str(sp_m), int(sl), float(tg), float(pr), float(ck), float(tt), float(sl + tg), "Lên đơn", int(ngay_len_don.month), int(ngay_len_don.year)], value_input_option="USER_ENTERED")
            st.session_state.order_items_count = 1; st.session_state.order_form_version += 1
            st.session_state.just_created_order = (ma_don, tot); refresh(); st.rerun()

elif nav == "💰 Lương Sale":
    banner("Doanh số & Lương Sale")
    if nhan_vien_df.empty or merged.empty: st.info("Chưa có đủ dữ liệu để tính lương.")
    else:
        ten_nv = st.selectbox("Chọn nhân viên Sale", nhan_vien_df["Ten_NV"].tolist())
        cm, cy = st.columns(2)
        thang = cm.selectbox("Tháng", list(range(1, 13)), index=(date.today().month - 1), key="sth")
        nam = cy.number_input("Năm", min_value=2020, max_value=2100, value=date.today().year, step=1, key="syr")
        nv_row = nhan_vien_df[nhan_vien_df["Ten_NV"] == ten_nv].iloc[0]
        ma_nv, ty_le_hh = nv_row["Ma_NV"], float(nv_row.get("Ty_le_hoa_hong") or 0.02)

        v_df = merged[merged["Trang_thai_Don"].isin(["Gửi kho", "Đang giao", "Đã giao"])].copy()
        v_df["Ngay_len_don"] = pd.to_datetime(v_df["Ngay_len_don"], errors="coerce")
        of_sale = v_df[(v_df["Sale_phu_trach"] == ma_nv) & (v_df["Ngay_len_don"].dt.month == thang) & (v_df["Ngay_len_don"].dt.year == nam)]
        dt = of_sale["Thanh_tien"].sum()
        dt_vat = dt * 0.92
        luong_cb = dt_vat * ty_le_hh

        is_dat = any(kw in str(ten_nv).lower() for kw in ["đạt", "dat"])
        ttan = tduc = 0
        if is_dat:
            tan_r = nhan_vien_df[nhan_vien_df["Ten_NV"].str.lower().str.contains("tân|tan", na=False)]
            if not tan_r.empty:
                ttan = v_df[(v_df["Sale_phu_trach"] == tan_r.iloc[0]["Ma_NV"]) & (v_df["Ngay_len_don"].dt.month == thang) & (v_df["Ngay_len_don"].dt.year == nam)]["Thanh_tien"].sum() * 0.92 * 0.01
            duc_r = nhan_vien_df[nhan_vien_df["Ten_NV"].str.lower().str.contains("đức|duc", na=False)]
            if not duc_r.empty:
                tduc = v_df[(v_df["Sale_phu_trach"] == duc_r.iloc[0]["Ma_NV"]) & (v_df["Ngay_len_don"].dt.month == thang) & (v_df["Ngay_len_don"].dt.year == nam)]["Thanh_tien"].sum() * 0.92 * 0.01

        luong_tong = luong_cb + ttan + tduc
        c1, c2 = st.columns(2)
        c1.markdown(f'<div class="metric-box"><div class="metric-label">Doanh thu sau VAT 8%</div><div class="metric-value" style="color:{GREEN};">{money(dt_vat)}</div></div>', unsafe_allow_html=True)
        c2.markdown(f'<div class="metric-box"><div class="metric-label">Lương thực nhận</div><div class="metric-value" style="color:{RED};">{money(luong_tong)}</div></div>', unsafe_allow_html=True)

        if not of_sale.empty:
            df_nhom = of_sale.merge(khach_hang_df[["Ma_KH", "Ten_NPP"]], on="Ma_KH", how="left").merge(san_pham_df[["Ma_SP", "Nhom_danh_muc"]], on="Ma_SP", how="left")
            by_npp = df_nhom.groupby(["Ten_NPP", "Nhom_danh_muc"]).agg(San_luong=("San_luong_xuat_kho", "sum"), Doanh_thu=("Thanh_tien", "sum")).reset_index()
            by_npp_disp = by_npp.copy(); by_npp_disp["Doanh_thu"] = by_npp_disp["Doanh_thu"].apply(money)
            st.dataframe(by_npp_disp, hide_index=True, use_container_width=True)

elif nav == "📊 Dashboard":
    banner("Tổng quan", highlight_text="Báo cáo & Đối soát kinh doanh")
    if merged.empty:
        st.info("Chưa có dữ liệu đơn hàng để phân tích.")
    else:
        # 1. XÁC ĐỊNH ĐƠN HÀNG HỢP LỆ THEO TIÊU CHUẨN THỰC THU & CÔNG NỢ
        # Ghi nhận doanh thu / công nợ cho các đơn:
        # - Đã gửi kho ("Gửi kho"), Đang giao ("Đang giao"), Đã giao ("Đã giao")
        # - HOẶC đơn đã thanh toán (bắt đầu bằng "TT -")
        def is_revenue_eligible(r):
            o_st = safe_str(r.get("Trang_thai_Don"))
            p_st = safe_str(r.get("Trang_thai_TT"))
            is_in_operation = o_st in ["Gửi kho", "Đang giao", "Đã giao"]
            is_paid = p_st.startswith("TT -") or ("đã vat" in p_st.lower() and "chưa" not in p_st.lower())
            return is_in_operation or is_paid

        def is_order_paid(p_st):
            s = safe_str(p_st)
            return s.startswith("TT -") or ("đã vat" in s.lower() and "chưa" not in s.lower())

        valid_orders = merged[merged.apply(is_revenue_eligible, axis=1)].copy()
        valid_orders["Ngay_len_don_dt"] = pd.to_datetime(valid_orders["Ngay_len_don"], errors="coerce")

        # Tạo danh sách bộ lọc thời gian
        time_dash_options = ["Tất cả thời gian"]
        valid_dates = valid_orders["Ngay_len_don_dt"].dropna()
        if not valid_dates.empty:
            for m in valid_dates.dt.to_period("M").drop_duplicates().sort_values(ascending=False):
                time_dash_options.append(f"Tháng {m.month:02d}/{m.year}")
            for y in valid_dates.dt.year.drop_duplicates().sort_values(ascending=False):
                if f"Năm {y}" not in time_dash_options:
                    time_dash_options.append(f"Năm {y}")

        # Danh sách Nhóm danh mục
        col_dm = next((c for c in san_pham_df.columns if "nhom" in str(c).lower() or "danh_muc" in str(c).lower()), "Nhom_danh_muc")
        list_dm = sorted([str(x).strip() for x in san_pham_df[col_dm].dropna().unique() if str(x).strip()]) if col_dm in san_pham_df.columns else []

        # Danh sách Sale
        list_sales = ["Tất cả Sale"]
        if "Ten_NV" in nhan_vien_df.columns:
            list_sales += sorted([str(x).strip() for x in nhan_vien_df["Ten_NV"].dropna().unique() if str(x).strip()])

        # Danh sách NPP
        list_npp = ["Tất cả NPP"]
        if "Ten_NPP" in khach_hang_df.columns:
            list_npp += sorted([str(x).strip() for x in khach_hang_df["Ten_NPP"].dropna().unique() if str(x).strip()])

        # Bố trí hàng lọc 5 cột
        f_c1, f_c2, f_c3, f_c4, f_c5 = st.columns([1.3, 1.3, 1.1, 1.3, 1.2])
        curr_m_str = f"Tháng {date.today().month:02d}/{date.today().year}"
        default_t_idx = time_dash_options.index(curr_m_str) if curr_m_str in time_dash_options else 0
        sel_time = f_c1.selectbox("Khoảng thời gian", time_dash_options, index=default_t_idx, key="dash_time")
        sel_dm = f_c2.multiselect("Nhóm danh mục", list_dm, default=[], placeholder="Tất cả nhóm hàng", key="dash_dm")
        sel_sale = f_c3.selectbox("Nhân viên Sale", list_sales, key="dash_sale")
        sel_npp = f_c4.selectbox("Nhà phân phối", list_npp, key="dash_npp")
        sel_pay = f_c5.selectbox("Dòng tiền", ["Tất cả dòng tiền", "Thực nhận (Đã TT)", "Chờ nhận (Công nợ)"], key="dash_pay")

        # 2. XỬ LÝ LỌC DỮ LIỆU
        filtered_df = valid_orders.copy()

        # Lọc theo thời gian
        if sel_time != "Tất cả thời gian":
            if sel_time.startswith("Tháng "):
                m_val, y_val = map(int, sel_time.replace("Tháng ", "").strip().split("/"))
                filtered_df = filtered_df[(filtered_df["Ngay_len_don_dt"].dt.month == m_val) & (filtered_df["Ngay_len_don_dt"].dt.year == y_val)]
            elif sel_time.startswith("Năm "):
                y_val = int(sel_time.replace("Năm ", "").strip())
                filtered_df = filtered_df[filtered_df["Ngay_len_don_dt"].dt.year == y_val]

        # Ghép thông tin sản phẩm và nhóm danh mục
        if not filtered_df.empty:
            cols_sp = [c for c in ["Ma_SP", "Ten_SP", col_dm, "Don_vi_tinh"] if c in san_pham_df.columns]
            filtered_df = filtered_df.merge(san_pham_df[cols_sp], on="Ma_SP", how="left")
            filtered_df[col_dm] = filtered_df[col_dm].fillna("Khác")
            if "Don_vi_tinh" not in filtered_df.columns:
                filtered_df["Don_vi_tinh"] = "Thùng"
            filtered_df["Don_vi_tinh"] = filtered_df["Don_vi_tinh"].fillna("Thùng")

        # Lọc theo nhóm danh mục
        if sel_dm and not filtered_df.empty:
            filtered_df = filtered_df[filtered_df[col_dm].isin(sel_dm)]

        # Lọc theo Sale
        if sel_sale != "Tất cả Sale" and not filtered_df.empty:
            sale_match = nhan_vien_df[nhan_vien_df["Ten_NV"] == sel_sale]
            if not sale_match.empty:
                ma_nv_val = str(sale_match.iloc[0]["Ma_NV"]).strip()
                filtered_df = filtered_df[filtered_df["Sale_phu_trach"].astype(str).str.strip() == ma_nv_val]

        # Lọc theo NPP
        if sel_npp != "Tất cả NPP" and not filtered_df.empty:
            kh_match = khach_hang_df[khach_hang_df["Ten_NPP"] == sel_npp]
            if not kh_match.empty:
                ma_kh_val = str(kh_match.iloc[0]["Ma_KH"]).strip()
                filtered_df = filtered_df[filtered_df["Ma_KH"].astype(str).str.strip() == ma_kh_val]

        # Lọc theo dòng tiền
        if sel_pay == "Thực nhận (Đã TT)" and not filtered_df.empty:
            filtered_df = filtered_df[filtered_df["Trang_thai_TT"].apply(is_order_paid)]
        elif sel_pay == "Chờ nhận (Công nợ)" and not filtered_df.empty:
            filtered_df = filtered_df[~filtered_df["Trang_thai_TT"].apply(is_order_paid)]

        st.write("")

        def fmt_qty(num):
            s = f"{float(num):,.1f}"
            if s.endswith(".0"):
                s = s[:-2]
            return s.replace(",", ".")

        # 3. TÍNH TOÁN BÓC TÁCH DÒNG TIỀN VÀ LƯỢNG HÀNG
        dt_tong_hop_le = filtered_df["Thanh_tien"].sum() if not filtered_df.empty else 0.0
        
        paid_mask = filtered_df["Trang_thai_TT"].apply(is_order_paid) if not filtered_df.empty else pd.Series(dtype=bool)
        paid_df = filtered_df[paid_mask] if not filtered_df.empty else pd.DataFrame()
        unpaid_df = filtered_df[~paid_mask] if not filtered_df.empty else pd.DataFrame()

        dt_thuc_nhan = paid_df["Thanh_tien"].sum() if not paid_df.empty else 0.0
        sl_thuc_nhan = (paid_df["SL_dat"].sum() + paid_df["Tang"].sum()) if not paid_df.empty else 0.0
        don_thuc_nhan = paid_df["Ma_don"].nunique() if not paid_df.empty else 0

        dt_cho_nhan = unpaid_df["Thanh_tien"].sum() if not unpaid_df.empty else 0.0
        sl_cho_nhan = (unpaid_df["SL_dat"].sum() + unpaid_df["Tang"].sum()) if not unpaid_df.empty else 0.0
        don_cho_nhan = unpaid_df["Ma_don"].nunique() if not unpaid_df.empty else 0
        
        dt_thuan_vat8 = dt_tong_hop_le * 0.92
        tong_sl_dat = filtered_df["SL_dat"].sum() if not filtered_df.empty else 0.0
        tong_sl_tang = filtered_df["Tang"].sum() if not filtered_df.empty else 0.0
        tong_sl_xuat = tong_sl_dat + tong_sl_tang
        so_don_hop_le = filtered_df["Ma_don"].nunique() if not filtered_df.empty else 0
        aov = (dt_tong_hop_le / so_don_hop_le) if so_don_hop_le > 0 else 0.0

        # HÀNG 4 THẺ CHỈ SỐ KPI CHÍNH
        k1, k2, k3, k4 = st.columns(4)
        with k1:
            st.markdown(f"""
            <div class="card-wrapper" style="background:#E2EDE8;border:2px solid #15503F;padding:12px 14px;border-radius:12px;">
                <div style="font-size:12px;font-weight:600;color:#15503F;">DOANH THU GHI NHẬN</div>
                <div style="font-size:19px;font-weight:800;color:#15503F;margin-top:2px;">{money(dt_tong_hop_le)}</div>
                <div style="font-size:11px;color:#4B5563;margin-top:2px;">Doanh thu thuần (-8% VAT): <b>{money(dt_thuan_vat8)}</b> ({so_don_hop_le} đơn)</div>
            </div>""", unsafe_allow_html=True)
        with k2:
            st.markdown(f"""
            <div class="card-wrapper" style="background:#DCFAF4;border:2px solid #0D9488;padding:12px 14px;border-radius:12px;">
                <div style="font-size:12px;font-weight:600;color:#0D9488;">THỰC NHẬN (ĐÃ TT)</div>
                <div style="font-size:19px;font-weight:800;color:#0D9488;margin-top:2px;">{money(dt_thuc_nhan)}</div>
                <div style="font-size:11px;color:#0D9488;margin-top:2px;">Lượng hàng: <b>{fmt_qty(sl_thuc_nhan)} thùng</b> ({don_thuc_nhan} đơn)</div>
            </div>""", unsafe_allow_html=True)
        with k3:
            st.markdown(f"""
            <div class="card-wrapper" style="background:#FDE8E8;border:2px solid #D92B2B;padding:12px 14px;border-radius:12px;">
                <div style="font-size:12px;font-weight:600;color:#D92B2B;">CHỜ NHẬN (CÔNG NỢ)</div>
                <div style="font-size:19px;font-weight:800;color:#D92B2B;margin-top:2px;">{money(dt_cho_nhan)}</div>
                <div style="font-size:11px;color:#D92B2B;margin-top:2px;">Hàng nợ: <b>{fmt_qty(sl_cho_nhan)} thùng</b> ({don_cho_nhan} đơn)</div>
            </div>""", unsafe_allow_html=True)
        with k4:
            st.markdown(f"""
            <div class="card-wrapper" style="background:#FEF3C7;border:2px solid #D97706;padding:12px 14px;border-radius:12px;">
                <div style="font-size:12px;font-weight:600;color:#D97706;">TỔNG HÀNG HÓA</div>
                <div style="font-size:19px;font-weight:800;color:#D97706;margin-top:2px;">{fmt_qty(tong_sl_xuat)} thùng</div>
                <div style="font-size:11px;color:#4B5563;margin-top:2px;">Đặt: <b>{fmt_qty(tong_sl_dat)}</b> | Tặng: <b>{fmt_qty(tong_sl_tang)}</b></div>
            </div>""", unsafe_allow_html=True)

        st.caption(f"📌 *Đang tổng hợp các đơn có tiến độ Gửi kho, Đang giao, Đã giao hoặc Đã TT. Giá trị đơn TB: **{money(aov)}**/đơn.*")
        st.write("")

        # BẢNG THEO DÕI ĐƠN HÀNG CÔNG NỢ & LƯỢNG HÀNG CHỜ THU (DÀNH RIÊNG ĐỂ KHÔNG BỎ SÓT)
        if not unpaid_df.empty:
            with st.expander(f"⚠️ **DANH SÁCH ĐƠN CHỜ GIAO & CHỜ THANH TOÁN ({don_cho_nhan} đơn - {fmt_qty(sl_cho_nhan)} thùng - {money(dt_cho_nhan)})**", expanded=True):
                unpaid_merged = unpaid_df.copy()
                unpaid_merged = unpaid_merged.merge(khach_hang_df[["Ma_KH", "Ten_NPP"]], on="Ma_KH", how="left")
                unpaid_merged["Ten_NPP"] = unpaid_merged["Ten_NPP"].fillna(unpaid_merged["Ma_KH"]).fillna("Chưa rõ NPP")
                
                # Group theo từng đơn để xem rõ
                debt_orders = unpaid_merged.groupby(["Ma_don", "Ngay_len_don", "Ten_NPP", "Trang_thai_Don", "Trang_thai_TT"]).agg(
                    SL_dat=("SL_dat", "sum"),
                    Tang=("Tang", "sum"),
                    Thanh_tien=("Thanh_tien", "sum")
                ).reset_index()
                debt_orders["Tong_thung"] = debt_orders["SL_dat"] + debt_orders["Tang"]
                debt_orders = debt_orders.sort_values("Thanh_tien", ascending=False)

                debt_disp = pd.DataFrame({
                    "Mã đơn": debt_orders["Ma_don"],
                    "Ngày lên đơn": debt_orders["Ngay_len_don"].astype(str),
                    "Nhà phân phối": debt_orders["Ten_NPP"],
                    "Tiến độ hàng": debt_orders["Trang_thai_Don"],
                    "Trạng thái TT": debt_orders["Trang_thai_TT"],
                    "SL thùng": debt_orders["Tong_thung"].apply(fmt_qty),
                    "Số tiền nợ (VNĐ)": debt_orders["Thanh_tien"].apply(money)
                })
                st.dataframe(debt_disp, hide_index=True, use_container_width=True)

        # 4. BẢNG CHI TIẾT: ĐỐI CHIẾU SẢN LƯỢNG SẢN PHẨM (NHÀ CUNG CẤP & KHO)
        st.markdown("<div style='font-size:14px;font-weight:700;color:#15503F;margin:10px 0 6px 0;'>📦 TỔNG HỢP LƯỢNG HÀNG XUẤT KHO</div>", unsafe_allow_html=True)
        
        if not filtered_df.empty:
            ncc_summary = filtered_df.groupby(["Ma_SP", "Ten_SP", col_dm, "Don_vi_tinh"]).agg(
                SL_dat=("SL_dat", "sum"),
                SL_tang=("Tang", "sum"),
                Doanh_thu=("Thanh_tien", "sum")
            ).reset_index()
            ncc_summary["Tong_xuat_kho"] = ncc_summary["SL_dat"] + ncc_summary["SL_tang"]
            ncc_summary = ncc_summary.sort_values("Tong_xuat_kho", ascending=False)

            # DataFrame hiển thị lên UI
            ncc_disp = pd.DataFrame({
                "Mã SP": ncc_summary["Ma_SP"],
                "Tên sản phẩm": ncc_summary["Ten_SP"],
                "Nhóm danh mục": ncc_summary[col_dm],
                "ĐVT": ncc_summary["Don_vi_tinh"],
                "SL đặt": ncc_summary["SL_dat"].apply(fmt_qty),
                "SL tặng": ncc_summary["SL_tang"].apply(fmt_qty),
                "Tổng xuất kho": ncc_summary["Tong_xuat_kho"].apply(fmt_qty),
                "Doanh thu sau CK": ncc_summary["Doanh_thu"].apply(money)
            })
            st.dataframe(ncc_disp, hide_index=True, use_container_width=True)

            # Nút tải Excel chuyên nghiệp cho bảng đối chiếu NCC
            exp_ncc_df = pd.DataFrame({
                "Mã SP": ncc_summary["Ma_SP"],
                "Tên sản phẩm": ncc_summary["Ten_SP"],
                "Nhóm danh mục": ncc_summary[col_dm],
                "Đơn vị tính": ncc_summary["Don_vi_tinh"],
                "Số lượng đặt": ncc_summary["SL_dat"],
                "Số lượng tặng": ncc_summary["SL_tang"],
                "Tổng sản lượng xuất kho": ncc_summary["Tong_xuat_kho"],
                "Doanh thu sau CK (VNĐ)": ncc_summary["Doanh_thu"]
            })
            b_ncc, m_ncc, ext_ncc = export_df_to_excel({"Doi_chieu_san_luong_NCC": exp_ncc_df})
            st.download_button("📥 Tải bảng đối chiếu NCC & Kho (Excel)", data=b_ncc, file_name=f"VIFEX_DoiChieu_NCC_{date.today().strftime('%Y%m%d')}.{ext_ncc}", mime=m_ncc, key="btn_dl_ncc_excel", use_container_width=True)
        else:
            st.info("Không có dữ liệu sản phẩm trong tiêu chí lọc đã chọn.")

        st.write("")

        # 5. PHÂN TÍCH CHUYÊN SÂU 2 CỘT (CƠ CẤU NHÓM HÀNG & TOP KHÁCH HÀNG NPP)
        col_an1, col_an2 = st.columns(2)
        with col_an1:
            st.markdown("<div style='font-size:13.5px;font-weight:700;color:#15503F;margin-bottom:6px;'>📊 CƠ CẤU & SẢN LƯỢNG THEO NHÓM HÀNG</div>", unsafe_allow_html=True)
            if not filtered_df.empty:
                by_cat = filtered_df.groupby(col_dm).agg(
                    SL_dat=("SL_dat", "sum"),
                    SL_tang=("Tang", "sum"),
                    Doanh_thu=("Thanh_tien", "sum")
                ).reset_index()
                by_cat["San_luong"] = by_cat["SL_dat"] + by_cat["SL_tang"]
                by_cat = by_cat.sort_values("Doanh_thu", ascending=False)
                tot_cat_dt = by_cat["Doanh_thu"].sum()
                by_cat["Ty_trong"] = by_cat["Doanh_thu"].apply(lambda v: f"{(v / tot_cat_dt * 100):.1f}%" if tot_cat_dt > 0 else "0%")
                
                cat_disp = pd.DataFrame({
                    "Nhóm hàng": by_cat[col_dm],
                    "Sản lượng (thùng)": by_cat["San_luong"].apply(fmt_qty),
                    "Doanh thu": by_cat["Doanh_thu"].apply(money),
                    "Tỷ trọng": by_cat["Ty_trong"]
                })
                st.dataframe(cat_disp, hide_index=True, use_container_width=True)
            else:
                st.info("Chưa có số liệu nhóm hàng.")

        with col_an2:
            st.markdown("<div style='font-size:13.5px;font-weight:700;color:#15503F;margin-bottom:6px;'>🏆 TOP 10 NHÀ PHÂN PHỐI NĂNG SUẤT</div>", unsafe_allow_html=True)
            if not filtered_df.empty:
                filtered_with_npp = filtered_df.merge(khach_hang_df[["Ma_KH", "Ten_NPP"]], on="Ma_KH", how="left")
                filtered_with_npp["Ten_NPP"] = filtered_with_npp["Ten_NPP"].fillna(filtered_with_npp["Ma_KH"]).fillna("Chưa rõ NPP")
                by_npp_dash = filtered_with_npp.groupby("Ten_NPP").agg(
                    So_don=("Ma_don", "nunique"),
                    SL_dat=("SL_dat", "sum"),
                    SL_tang=("Tang", "sum"),
                    Doanh_thu=("Thanh_tien", "sum")
                ).reset_index()
                by_npp_dash["San_luong"] = by_npp_dash["SL_dat"] + by_npp_dash["SL_tang"]
                by_npp_dash = by_npp_dash.sort_values("Doanh_thu", ascending=False).head(10)

                npp_disp = pd.DataFrame({
                    "Nhà phân phối": by_npp_dash["Ten_NPP"],
                    "Số đơn": by_npp_dash["So_don"],
                    "Sản lượng (thùng)": by_npp_dash["San_luong"].apply(fmt_qty),
                    "Doanh số": by_npp_dash["Doanh_thu"].apply(money)
                })
                st.dataframe(npp_disp, hide_index=True, use_container_width=True)
            else:
                st.info("Chưa có số liệu khách hàng.")

        st.write("")

        # 6. HIỆU SUẤT ĐỘI NGŨ SALE
        st.markdown("<div style='font-size:13.5px;font-weight:700;color:#15503F;margin:10px 0 6px 0;'>👥 HIỆU SUẤT BÁN HÀNG</div>", unsafe_allow_html=True)
        if not filtered_df.empty:
            sale_merged = filtered_df.copy()
            if "Ma_NV" in nhan_vien_df.columns and "Ten_NV" in nhan_vien_df.columns:
                sale_merged = sale_merged.merge(nhan_vien_df[["Ma_NV", "Ten_NV"]], left_on="Sale_phu_trach", right_on="Ma_NV", how="left")
                sale_merged["Ten_NV"] = sale_merged["Ten_NV"].fillna(sale_merged["Sale_phu_trach"]).fillna("Chưa rõ Sale")
            else:
                sale_merged["Ten_NV"] = sale_merged["Sale_phu_trach"].fillna("Chưa rõ Sale")

            by_sale_dash = sale_merged.groupby("Ten_NV").agg(
                So_don=("Ma_don", "nunique"),
                SL_dat=("SL_dat", "sum"),
                SL_tang=("Tang", "sum"),
                Doanh_thu=("Thanh_tien", "sum")
            ).reset_index()
            by_sale_dash["San_luong"] = by_sale_dash["SL_dat"] + by_sale_dash["SL_tang"]
            by_sale_dash = by_sale_dash.sort_values("Doanh_thu", ascending=False)
            by_sale_dash["Doanh_thu_vat8"] = by_sale_dash["Doanh_thu"] * 0.92

            sale_disp = pd.DataFrame({
                "Nhân viên Sale": by_sale_dash["Ten_NV"],
                "Số đơn": by_sale_dash["So_don"],
                "Sản lượng (thùng)": by_sale_dash["San_luong"].apply(fmt_qty),
                "Doanh thu hợp lệ": by_sale_dash["Doanh_thu"].apply(money),
                "Doanh thu thuần (-8% VAT)": by_sale_dash["Doanh_thu_vat8"].apply(money)
            })
            st.dataframe(sale_disp, hide_index=True, use_container_width=True)
        else:
            st.info("Chưa có số liệu nhân viên Sale.")

elif nav == "👥 Khách hàng":
    banner("Danh sách khách hàng")
    cols = [c for c in ["Ma_KH", "Ten_NPP", "Sale_phu_trach", "Khu_vuc", "Trang_thai"] if c in khach_hang_df.columns]
    show = khach_hang_df[cols].dropna(subset=["Ten_NPP"]).copy() if "Ten_NPP" in khach_hang_df.columns else khach_hang_df.copy()
    st.dataframe(show, hide_index=True, use_container_width=True)
